"""按表导出到 Excel（可带当前筛选条件）。"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from database import (
    INDEX_LABEL,
    PRODUCT_FIELDS,
    QTY_LABEL,
    STOCK_LABEL,
    SUMMARY_PRICE_LABEL,
    format_number,
)


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws):
    for column in ws.columns:
        letter = column[0].column_letter
        width = 10
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            width = max(width, min(len(value) + 4, 40))
        ws.column_dimensions[letter].width = width


def _save_sheet(path, title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header(ws)
    _autosize(ws)
    wb.save(path)
    return path


def export_products(db, path, keyword=""):
    product_customs = db.custom_fields("product")
    headers = [INDEX_LABEL] + [label for _, label, _, _ in PRODUCT_FIELDS] + [
        f["label"] for f in product_customs
    ]
    rows = []
    for idx, item in enumerate(db.list_products(keyword), start=1):
        row = [idx]
        for key, _, data_type, _ in PRODUCT_FIELDS:
            value = item.get(key)
            row.append(format_number(value) if data_type == "number" else (value or ""))
        for field in product_customs:
            row.append(db.extra_text(item.get("extra"), field))
        rows.append(row)
    return _save_sheet(path, "商品基础信息表", headers, rows)


def export_transactions(db, path, date_from="", date_to="", io_type="", keyword=""):
    product_customs = db.custom_fields("product")
    txn_customs = db.custom_fields("txn")
    headers = [
        INDEX_LABEL,
        "日期",
        "类型（入库/出库）",
        "产品型号",
        "颜色",
        "单价",
        QTY_LABEL,
        "备注",
    ]
    headers += [f["label"] for f in product_customs]
    headers += [f["label"] for f in txn_customs]
    rows = []
    for idx, item in enumerate(
        db.list_transactions(date_from, date_to, io_type, keyword), start=1
    ):
        row = [
            idx,
            item["txn_date"],
            item["io_type"],
            item.get("product_model") or "",
            item["color"] or "",
            format_number(item["unit_price"]),
            format_number(item["quantity"]),
            item["remark"] or "",
        ]
        for field in product_customs:
            row.append(db.extra_text(item.get("extra"), field))
        for field in txn_customs:
            row.append(db.extra_text(item.get("extra"), field))
        rows.append(row)
    return _save_sheet(path, "出入库流水表", headers, rows)


def export_summary(db, path, hide_zero=False, keyword=""):
    product_customs = db.custom_fields("product")
    headers = [INDEX_LABEL, "产品型号", "颜色", STOCK_LABEL, SUMMARY_PRICE_LABEL]
    headers += [f["label"] for f in product_customs]
    rows = []
    for idx, item in enumerate(db.summary(hide_zero=hide_zero, keyword=keyword), start=1):
        row = [
            idx,
            item.get("product_model") or "",
            item["color"] or "",
            format_number(item["stock"]),
            format_number(item.get("unit_price")),
        ]
        for field in product_customs:
            row.append(db.extra_text(item.get("extra"), field))
        rows.append(row)
    return _save_sheet(path, "库存汇总表", headers, rows)


def export_excel(db, path):
    """导出全部三张表（不做筛选，便于整库备份）。"""
    product_customs = db.custom_fields("product")
    txn_customs = db.custom_fields("txn")
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "商品基础信息表"
    headers1 = [INDEX_LABEL] + [label for _, label, _, _ in PRODUCT_FIELDS] + [
        f["label"] for f in product_customs
    ]
    ws1.append(headers1)
    for idx, item in enumerate(db.list_products(), start=1):
        row = [idx]
        for key, _, data_type, _ in PRODUCT_FIELDS:
            value = item.get(key)
            row.append(format_number(value) if data_type == "number" else (value or ""))
        for field in product_customs:
            row.append(db.extra_text(item.get("extra"), field))
        ws1.append(row)
    _style_header(ws1)
    _autosize(ws1)

    ws2 = wb.create_sheet("出入库流水表")
    headers2 = [
        INDEX_LABEL,
        "日期",
        "类型（入库/出库）",
        "产品型号",
        "颜色",
        "单价",
        QTY_LABEL,
        "备注",
    ]
    headers2 += [f["label"] for f in product_customs]
    headers2 += [f["label"] for f in txn_customs]
    ws2.append(headers2)
    for idx, item in enumerate(db.list_transactions(), start=1):
        row = [
            idx,
            item["txn_date"],
            item["io_type"],
            item.get("product_model") or "",
            item["color"] or "",
            format_number(item["unit_price"]),
            format_number(item["quantity"]),
            item["remark"] or "",
        ]
        for field in product_customs:
            row.append(db.extra_text(item.get("extra"), field))
        for field in txn_customs:
            row.append(db.extra_text(item.get("extra"), field))
        ws2.append(row)
    _style_header(ws2)
    _autosize(ws2)

    ws3 = wb.create_sheet("库存汇总表")
    headers3 = [INDEX_LABEL, "产品型号", "颜色", STOCK_LABEL, SUMMARY_PRICE_LABEL]
    headers3 += [f["label"] for f in product_customs]
    ws3.append(headers3)
    for idx, item in enumerate(db.summary(hide_zero=False), start=1):
        row = [
            idx,
            item.get("product_model") or "",
            item["color"] or "",
            format_number(item["stock"]),
            format_number(item.get("unit_price")),
        ]
        for field in product_customs:
            row.append(db.extra_text(item.get("extra"), field))
        ws3.append(row)
    _style_header(ws3)
    _autosize(ws3)

    wb.save(path)
    return path
